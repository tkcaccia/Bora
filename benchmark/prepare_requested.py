"""Inventory SICAPv2 and explicitly bounded CAMELYON16 / SegPath acquisitions."""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import numpy as np
from PIL import Image

from .download import DEFAULT_ROOT, LOCK, digest
from .prepare import entry, write_once


def provider_tables(raw, root):
    from openpyxl import load_workbook
    output = {}
    for path in sorted(raw.rglob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=True)
        if len(book.sheetnames) != 1:
            raise ValueError("Unexpected provider table structure")
        data = iter(book.active.values)
        header = next(data)
        output[str(path.relative_to(raw))] = [dict(zip(header, values)) for values in data]
        book.close()
    proof = {"tables": output, "sha256": {str(p.relative_to(raw)): digest(p) for p in raw.rglob("*.xlsx")}}
    (root / "metadata" / "sicapv2-table-inventory.json").write_text(json.dumps(proof, indent=2) + "\n")
    return output


def sicapv2(root):
    raw = root / "raw" / "sicapv2" / "extracted" / "SICAPv2"
    tables = provider_tables(raw, root)
    slides = {str(r["slide_id"]): r for r in tables["wsi_labels.xlsx"]}
    partitions, image_labels = {}, {}
    for split in ("Train", "Test"):
        for record in tables[f"partition/Test/{split}.xlsx"]:
            name = record["image_name"]
            if name in partitions:
                raise ValueError("Duplicate primary partition membership")
            partitions[name], image_labels[name] = split.lower(), record
    images = sorted((raw / "images").glob("*.jpg"))
    masks = sorted((raw / "masks").glob("*.jpg"))
    if len(images) != 18783 or {p.name for p in images} != {p.name for p in masks}:
        raise ValueError("Expected all 18,783 original JPEG pairs")
    rows = []
    for i, image in enumerate(images):
        slide = image.name.split("_", 1)[0]
        if slide not in slides:
            raise ValueError("Missing supplied patient mapping")
        record = slides[slide]
        row = entry(image, raw / "masks" / image.name, root, "sicapv2", image.stem,
                    "sicap-patient--" + str(record["patient_id"]),
                    partitions.get(image.name, "outside_released_classification_partition"),
                    "raw_jpeg_annotation_unreconciled", [], None)
        position = re.search(r"_xini_(\d+)_yini_(\d+)$", image.stem)
        row.update(slide_id=slide, group_level="provider_patient_id", xini=int(position[1]), yini=int(position[2]),
                   source_status="original_JPEG_masks_semantic_encoding_and_boundary_fidelity_unresolved",
                   gleason_primary=record["Gleason_primary"], gleason_secondary=record["Gleason_secondary"],
                   overlay_supported=False)
        if image.name in image_labels:
            row["grade"] = ", ".join(k for k, v in image_labels[image.name].items() if k != "image_name" and v == 1)
        rows.append(row)
        if (i + 1) % 3000 == 0:
            print(f"SICAPv2: inventoried {i + 1} original pairs", flush=True)
    groups = {split: {r["group_id"] for r in rows if r["split"] == split} for split in ("train", "test")}
    return rows, {"patients": len({r["group_id"] for r in rows}), "source_slide_table_rows": len(slides),
                  "image_slides": len({r["slide_id"] for r in rows}),
                  "patient_overlap_primary_train_test": sorted(groups["train"] & groups["test"]),
                  "use_status": "Complete original archive; keep out of primary boundary evaluation until lossless grade encoding is reconciled",
                  "source_mask_format": "JPEG grayscale; decoded intensities are not the documented conceptual 0-3 labels"}


def camelyon16_subset(root):
    import openslide
    raw = root / "raw" / "camelyon16-subset"
    source = json.loads(LOCK.read_text())["camelyon16-subset"]
    rows, full = [], []
    for sample in source["selection"]:
        image, mask, xml = raw / "images" / (sample + ".tif"), raw / "masks" / (sample + "_mask.tif"), raw / "annotations" / (sample + ".xml")
        with openslide.OpenSlide(image) as slide, openslide.OpenSlide(mask) as annotation:
            if slide.dimensions != annotation.dimensions:
                raise ValueError("WSI and mask native dimensions disagree")
            x_mpp = float(slide.properties["openslide.mpp-x"])
            y_mpp = float(slide.properties["openslide.mpp-y"])
            points = []
            for region in ET.parse(xml).getroot().findall(".//Annotation"):
                if region.get("PartOfGroup", "").lower() != "tumor":
                    continue
                vertices = region.findall(".//Coordinate")
                points.extend(vertices[::max(1, len(vertices) // 20)])
            chosen = None
            side = 4096
            for point in points:
                x = max(0, min(slide.dimensions[0] - side, round(float(point.get("X"))) - side // 2))
                y = max(0, min(slide.dimensions[1] - side, round(float(point.get("Y"))) - side // 2))
                values = np.asarray(annotation.read_region((x, y), 0, (side, side)))
                if not (np.array_equal(values[:, :, 0], values[:, :, 1]) and np.array_equal(values[:, :, 0], values[:, :, 2])):
                    raise ValueError("Expected equal channels for categorical mask")
                labels = values[:, :, 0]
                if not set(np.unique(labels)).issubset({0, 1, 2}):
                    raise ValueError("Unexpected native CAMELYON mask values")
                if .01 < np.mean(labels == 2) < .99 and np.mean(labels == 1) > .01:
                    chosen = x, y, labels.copy()
                    break
            if chosen is None:
                raise ValueError("No visible annotated tumor interface found in the bounded candidate set")
            x, y, labels = chosen
            out = root / "prepared" / "camelyon16-subset" / "native_examples"
            out.mkdir(parents=True, exist_ok=True)
            rgb, mk = out / (sample + ".png"), out / (sample + "_mask.png")
            from io import BytesIO
            for payload, path in ((slide.read_region((x, y), 0, (side, side)).convert("RGB"), rgb),
                                  (Image.fromarray(labels), mk)):
                buffer = BytesIO()
                payload.save(buffer, format="PNG")
                write_once(path, buffer.getvalue())
            row = entry(rgb, mk, root, "camelyon16-subset", sample + "--native-example", sample,
                        "train", "semantic", [0], None)
            row.update(slide_id=sample, group_level="slide_patient_mapping_unavailable", mpp_x_um=x_mpp, mpp_y_um=y_mpp,
                       resolution_status="embedded_WSI_axis_spacing", source_status="original_publisher_WSI_and_mask_native_crop",
                       crop_x=x, crop_y=y, crop_level=0, crop_width=side, crop_height=side,
                       preprocessing="Level-0 example window only; no resampling or new mask rasterization")
            rows.append(row)
            full.append({"slide": sample, "image": str(image.relative_to(root)), "mask": str(mask.relative_to(root)),
                         "xml": str(xml.relative_to(root)), "image_sha256": digest(image), "mask_sha256": digest(mask),
                         "xml_sha256": digest(xml), "width": slide.dimensions[0], "height": slide.dimensions[1],
                         "mpp_x_um": x_mpp, "mpp_y_um": y_mpp, "image_levels": slide.level_dimensions,
                         "mask_levels": annotation.level_dimensions, "scanner_vendor": slide.properties.get("openslide.vendor"),
                         "example_window": {k: row[k] for k in ("crop_x", "crop_y", "crop_level", "crop_width", "crop_height")}})
    (root / "metadata" / "camelyon16-subset-wsi.json").write_text(json.dumps(full, indent=2) + "\n")
    return rows, {"complete_WSIs": len(full), "native_example_windows": len(rows),
                  "scope": "Three positive training WSIs; convenience windows are not prespecified evaluation sampling",
                  "use_status": "Tumor annotation is expert-derived; normal tissue class is automatically generated, not expert stroma"}


def segpath_subset(root):
    source = json.loads(LOCK.read_text())["segpath-subset"]
    rows, original_counts = [], {}
    for target in source["targets"]:
        receipt = root / "raw" / "segpath-subset" / target["class"] / "subset-receipt.json"
        info = json.loads(receipt.read_text())
        table = next(x for x in target["sidecars"] if x["path"].endswith(".csv"))
        with (root / "raw" / "segpath-subset" / table["path"]).open() as handle:
            splits = {Path(r["filename"]).name.removesuffix("_mask.png").removesuffix("_HE.png"): r["train_val_test"] for r in csv.DictReader(handle)}
        original_counts[target["class"]] = {"patches_listed_in_complete_CSV": len(splits),
                                             "released_partitions": dict(Counter(splits.values()))}
        for pair in info["pairs"]:
            dataset = "segpath-" + target["class"].split("_", 1)[0].lower()
            row = entry(root / pair["image"], root / pair["mask"], root, dataset, pair["id"],
                        "segpath-source-slide--" + pair["slide_id"], splits[pair["id"]], "binary_semantic", [], 0.220818)
            row.update(source_status="original_TAR_native_subset_full_archive_checksum_unverified",
                       group_level="source_slide_id_TMA_spot_patient_mapping_unavailable", target=target["class"],
                       source_record=target["record"], subset_role="convenience_inspection_only")
            rows.append(row)
    return rows, {"targets": len(source["targets"]), "scope": "Three native pairs per target; not a representative benchmark cohort",
                  "full_archive_checksum_verified": False, "original_CSV_cohorts": original_counts}


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("dataset", choices=("sicapv2", "camelyon16-subset", "segpath-subset"))
    p.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    a = p.parse_args(argv)
    root = a.root.resolve()
    rows, extra = globals()[a.dataset.replace("-", "_")](root)
    inventories = {}
    for dataset in sorted({r["dataset"] for r in rows}):
        subset = sorted((r for r in rows if r["dataset"] == dataset), key=lambda r: r["id"])
        path = root / "prepared" / (dataset + ".jsonl")
        write_once(path, ("\n".join(json.dumps(r, sort_keys=True) for r in subset) + "\n").encode())
        inventories[dataset] = {"pairs": len(subset), "sha256": digest(path), "splits": dict(Counter(r["split"] for r in subset))}
    report = {"dataset": a.dataset, "inventories": inventories, **extra, "bora_executed": False}
    (root / "metadata" / (a.dataset + "-prepared.json")).write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
