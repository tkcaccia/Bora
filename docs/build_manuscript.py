#!/usr/bin/env python3
"""Build the Bora manuscript DOCX from its authoritative Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscript.md"
OUTPUT = ROOT / "Bora_manuscript.docx"
BLUE, DARK, GRAY = RGBColor(46,116,181), RGBColor(31,77,120), RGBColor(92,99,112)


def font(run, size=11, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end]); font(run, 9, color=GRAY)


def clean(text):
    text = text.replace("–", "-").replace("—", "-").replace("‑", "-").replace("→", "->")
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    replacements = {r"\(":"",r"\)":"",r"\operatorname":"",r"\mathrm":"",r"\text":"",
        r"\cup":" union ",r"\bigcup":"union",r"\setminus":" minus ",r"\qquad":"; ",
        r"\lVert":"||",r"\rVert":"||",r"\exp":"exp",r"\propto":" proportional to ",
        r"\beta":"beta",r"\lambda":"lambda",r"\sum":"sum",r"\sqrt{2}":"sqrt(2)",
        r"\in":" in ",r"\rightarrow":" -> ",r"\Omega":"Omega"}
    for old,new in replacements.items(): text=text.replace(old,new)
    text = text.replace("{", "").replace("}", "").replace("\\", "")
    return text.replace("**", "").replace("`", "").replace("*", "")


def new_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids or [0])+1, max(num_ids or [0])+1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"),str(abstract_id))
    multi=OxmlElement("w:multiLevelType"); multi.set(qn("w:val"),"singleLevel"); abstract.append(multi)
    level=OxmlElement("w:lvl"); level.set(qn("w:ilvl"),"0")
    start=OxmlElement("w:start"); start.set(qn("w:val"),"1"); level.append(start)
    fmt=OxmlElement("w:numFmt"); fmt.set(qn("w:val"),"decimal"); level.append(fmt)
    text=OxmlElement("w:lvlText"); text.set(qn("w:val"),"%1."); level.append(text)
    ppr=OxmlElement("w:pPr"); tabs=OxmlElement("w:tabs"); tab=OxmlElement("w:tab")
    tab.set(qn("w:val"),"num"); tab.set(qn("w:pos"),"540"); tabs.append(tab); ppr.append(tabs)
    ind=OxmlElement("w:ind"); ind.set(qn("w:left"),"540"); ind.set(qn("w:hanging"),"280"); ppr.append(ind)
    level.append(ppr); abstract.append(level); numbering.append(abstract)
    num=OxmlElement("w:num"); num.set(qn("w:numId"),str(num_id))
    ref=OxmlElement("w:abstractNumId"); ref.set(qn("w:val"),str(abstract_id)); num.append(ref); numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr=paragraph._p.get_or_add_pPr(); numpr=OxmlElement("w:numPr")
    ilvl=OxmlElement("w:ilvl"); ilvl.set(qn("w:val"),"0")
    nid=OxmlElement("w:numId"); nid.set(qn("w:val"),str(num_id))
    numpr.extend([ilvl,nid]); ppr.append(numpr)


def add_inline(paragraph, text):
    text = clean(text)
    cursor = 0
    for match in re.finditer(r"(\[[A-Z][^]]*\])", text):
        if match.start() > cursor: font(paragraph.add_run(text[cursor:match.start()]))
        run = paragraph.add_run(match.group(1)); font(run, bold=True, color=RGBColor(156,92,0))
        run.font.highlight_color = 7
        cursor = match.end()
    if cursor < len(text): font(paragraph.add_run(text[cursor:]))


doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
for name, size, color, before, after in (
    ("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)):
    style=doc.styles[name]; style.font.name="Calibri"; style.font.size=Pt(size); style.font.bold=True; style.font.color.rgb=color
    style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after); style.paragraph_format.keep_with_next=True
for name in ("List Bullet","List Number"):
    style=doc.styles[name]; style.font.name="Calibri"; style.font.size=Pt(11)
    style.paragraph_format.left_indent=Inches(.375); style.paragraph_format.first_line_indent=Inches(-.194)
    style.paragraph_format.space_after=Pt(4); style.paragraph_format.line_spacing=1.208

header=section.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.LEFT
font(header.add_run("Bora | Medical Image Analysis manuscript"),9,bold=True,color=GRAY)
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run("Protocol-stage draft  |  "),9,color=GRAY); add_page_field(footer)

lines=SOURCE.read_text().splitlines(); i=0; title_done=False; number_id=None; previous_numbered=False
while i < len(lines):
    raw=lines[i].rstrip(); stripped=raw.strip()
    if not stripped: i+=1; continue
    if stripped.startswith("# ") and not title_done:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(72); p.paragraph_format.space_after=Pt(14)
        font(p.add_run(clean(stripped[2:])),24,bold=True,color=DARK)
        title_done=True; i+=1; continue
    if stripped.startswith("### "): p=doc.add_paragraph(clean(stripped[4:]),style="Heading 3")
    elif stripped.startswith("## "): p=doc.add_paragraph(clean(stripped[3:]),style="Heading 2")
    elif stripped.startswith("# "): p=doc.add_paragraph(clean(stripped[2:]),style="Heading 1")
    elif stripped.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); add_inline(p,stripped[2:])
    elif re.match(r"^\d+\. ",stripped):
        if not previous_numbered: number_id=new_numbering(doc)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.208
        apply_numbering(p,number_id); add_inline(p,re.sub(r"^\d+\. ","",stripped)); previous_numbered=True
    elif stripped.startswith("\\["):
        equation=[stripped]; i+=1
        while i<len(lines) and not lines[i].strip().endswith("\\]"): equation.append(lines[i].strip()); i+=1
        if i<len(lines): equation.append(lines[i].strip())
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
        font(p.add_run(clean(" ".join(equation).replace("\\[","").replace("\\]",""))),10,italic=True,name="Cambria Math")
    else:
        chunks=[stripped]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |- |\d+\. |\\\[)",lines[i].strip()):
            chunks.append(lines[i].strip()); i+=1
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; add_inline(p," ".join(chunks)); previous_numbered=False; continue
    if not re.match(r"^\d+\. ",stripped): previous_numbered=False
    i+=1

doc.core_properties.title = "Bora: scalable boundary-constrained refinement of tile-level tissue labels"
doc.core_properties.subject = "Protocol-stage manuscript for Medical Image Analysis"
doc.core_properties.author = "Bora project contributors"
doc.save(OUTPUT)
print(OUTPUT)
